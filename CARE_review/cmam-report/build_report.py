"""Build the narrative CMAM report from its Markdown source."""
from pathlib import Path
import argparse
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def inline(paragraph, text):
    for token in re.split(r'(\*\*.*?\*\*|\*[^*]+\*)', text):
        if not token:
            continue
        run = paragraph.add_run(token.strip('*') if token.startswith('*') else token)
        run.bold = token.startswith('**')
        run.italic = token.startswith('*') and not token.startswith('**')


def build(source, output):
    raw = Path(source).read_text()
    if '\u2014' in raw:
        raise ValueError('The report must not contain em dashes')
    front, body = raw.split('---', 2)[1:]
    meta = dict(re.findall(r'^(\w+):\s*"(.*?)"\s*$', front, re.M))
    document = Document()
    sec = document.sections[0]
    sec.page_width, sec.page_height = Inches(8.2677), Inches(11.6929)
    sec.top_margin = sec.bottom_margin = Inches(.78)
    sec.left_margin = sec.right_margin = Inches(.84)
    sec.footer_distance = Inches(.35)
    sec.header_distance = Inches(.30)
    styles = document.styles
    # Remove decorative residue from the bundled default Word template.
    for style in styles:
        for e in list(style.element.iter()):
            if e.tag in (qn('w:pBdr'), qn('w:contextualSpacing')):
                e.getparent().remove(e)
            elif e.tag == qn('w:spacing') and e.getparent().tag == qn('w:rPr'):
                e.getparent().remove(e)
    for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption'):
        style = styles[name]
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor(0, 0, 0)
        rp = style.element.get_or_add_rPr()
        for e in list(rp):
            if e.tag == qn('w:color'):
                for attr in ('themeColor', 'themeTint', 'themeShade'):
                    e.attrib.pop(qn('w:'+attr), None)
        style.paragraph_format.widow_control = True
    normal = styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(7)
    styles['Title'].font.size = Pt(25)
    styles['Title'].font.bold = True
    styles['Title'].paragraph_format.space_after = Pt(9)
    styles['Title'].paragraph_format.keep_with_next = True
    styles['Subtitle'].font.size = Pt(12)
    styles['Subtitle'].font.italic = False
    styles['Subtitle'].paragraph_format.space_after = Pt(10)
    styles['Subtitle'].paragraph_format.keep_with_next = True
    for name, size, before, after in [('Heading 1',16,17,8),('Heading 2',12.5,12,6)]:
        st=styles[name]
        st.font.size=Pt(size)
        st.font.bold=True
        st.paragraph_format.space_before=Pt(before)
        st.paragraph_format.space_after=Pt(after)
        st.paragraph_format.keep_with_next=True
    styles['Caption'].font.size = Pt(10)
    styles['Caption'].font.italic = False
    styles['Caption'].font.bold = True
    styles['Caption'].paragraph_format.keep_with_next = True
    document.add_paragraph(meta['title'], 'Title')
    document.add_paragraph(meta['subtitle'], 'Subtitle')
    p=document.add_paragraph(meta['date'])
    p.paragraph_format.space_after=Pt(16)
    p.runs[0].font.size=Pt(10)
    p.paragraph_format.keep_with_next=True
    references=False
    after_table=False
    for block in re.split(r'\n\s*\n',body.strip()):
        block=block.strip()
        if not block: continue
        if block.startswith('# '):
            heading=block[2:]
            p=document.add_paragraph(heading,'Heading 1')
            if heading=='References':
                references=True
                p.paragraph_format.page_break_before=True
            continue
        if block.startswith('## '):
            document.add_paragraph(block[3:],'Heading 2')
            continue
        if block.startswith('|'):
            rows=[[c.strip() for c in line.strip('|').split('|')] for line in block.splitlines() if line.startswith('|')]
            rows.pop(1)
            table=document.add_table(rows=len(rows),cols=len(rows[0]))
            table.alignment=WD_TABLE_ALIGNMENT.CENTER
            table.autofit=False
            widths=[2.10,1.10,1.60,1.79]
            for col,w in zip(table.columns,widths): col.width=Inches(w)
            for i,row in enumerate(table.rows):
                pr=row._tr.get_or_add_trPr()
                pr.append(OxmlElement('w:cantSplit'))
                if i==0: pr.append(OxmlElement('w:tblHeader'))
                for j,cell in enumerate(row.cells):
                    cell.width=Inches(widths[j])
                    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    tcpr=cell._tc.get_or_add_tcPr()
                    borders=OxmlElement('w:tcBorders')
                    for edge in ('top','left','bottom','right'):
                        b=OxmlElement('w:'+edge)
                        for a,v in [('val','single'),('sz','4'),('color','D9D9D9')]: b.set(qn('w:'+a),v)
                        borders.append(b)
                    tcpr.append(borders)
                    margins=OxmlElement('w:tcMar')
                    for edge,amount in [('top',100),('bottom',100),('left',105),('right',105)]:
                        v=OxmlElement('w:'+edge);v.set(qn('w:w'),str(amount));v.set(qn('w:type'),'dxa');margins.append(v)
                    tcpr.append(margins)
                    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E6EBEE' if i==0 else 'FFFFFF');tcpr.append(sh)
                    p=cell.paragraphs[0]
                    p.paragraph_format.space_after=Pt(0)
                    p.paragraph_format.line_spacing=1.08
                    p.paragraph_format.keep_with_next=(i==0)
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                    inline(p,rows[i][j])
                    for r in p.runs: r.font.size=Pt(9.5);r.bold=(i==0)
            after_table=True
            continue
        if references:
            for ref in block.splitlines():
                p=document.add_paragraph()
                inline(p,ref.removeprefix('- '))
                p.paragraph_format.left_indent=Inches(.18)
                p.paragraph_format.first_line_indent=Inches(-.18)
                p.paragraph_format.keep_together=True
                p.paragraph_format.space_after=Pt(6)
                p.paragraph_format.line_spacing=1.08
                for r in p.runs: r.font.size=Pt(9.5)
        else:
            p=document.add_paragraph()
            inline(p,' '.join(block.splitlines()))
            if after_table:
                p.paragraph_format.space_before=Pt(9)
                after_table=False
            if block.startswith('**') and block.endswith('**'):
                p.style='Caption'
            if block.startswith('*To be completed'):
                p.paragraph_format.space_before=Pt(3)
    footer=sec.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=footer.add_run()
    r.font.name='Arial';r.font.size=Pt(9)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
    footer._p.append(field)
    document.core_properties.title=meta['title']
    document.core_properties.subject=meta['subtitle']
    document.core_properties.author=''
    document.core_properties.last_modified_by=''
    document.core_properties.comments=''
    document.save(output)
    print(f'Built {output}: {len(document.paragraphs)} paragraphs and {len(document.tables)} table')


if __name__=='__main__':
    here=Path(__file__).resolve().parent
    parser=argparse.ArgumentParser()
    parser.add_argument('--source',default=here/'CMAM_REPORT.md',type=Path)
    parser.add_argument('--output',default=here/'CMAM_REPORT.docx',type=Path)
    args=parser.parse_args()
    build(args.source,args.output)
